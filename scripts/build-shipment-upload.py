#!/usr/bin/env python3
"""Build a Supplier Hub shipment XLSX from a minimal PO/SKU/carton manifest."""

from __future__ import annotations

import argparse
import json
import re
import sys
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet


HEADER_PATTERNS = {
    "order_no": ("발주번호", "poid"),
    "center_code": ("물류센터", "fc"),
    "transport_type": ("입고유형", "transporttype"),
    "expected_inbound_date": ("입고예정일", "edd"),
    "sku_code": ("상품번호", "skuid", "skucode"),
    "barcode": ("상품바코드", "skubarcode", "barcode"),
    "sku_name": ("상품이름", "skuname"),
    "confirmed_quantity": ("확정수량", "confirmedqty", "confirmedquantity"),
    "slip_no": ("송장번호", "invoicenumber", "slipno"),
    "quantity": ("납품수량", "출고수량", "shippedqty", "shippedquantity"),
}


@dataclass(frozen=True)
class RowTemplate:
    cells: tuple[dict[str, Any], ...]
    height: float | None
    hidden: bool
    outline_level: int


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    source = Path(args.source).resolve()
    output = Path(args.output).resolve()
    if source == output:
        raise ValueError("원본 통합문서는 출력 파일로 덮어쓸 수 없습니다.")
    if output.exists():
        raise FileExistsError(f"출력 파일이 이미 존재합니다: {output}")

    manifest = json.load(sys.stdin)
    if set(manifest) != {"orders", "batches", "cartons"}:
        raise ValueError("manifest에는 orders, batches, cartons만 있어야 합니다.")

    workbook = load_workbook(source)
    worksheet, header_row, columns = find_product_sheet(workbook.worksheets)
    write_carton_rows(worksheet, header_row, columns, manifest)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()
    # ASCII escaping keeps Windows paths lossless even when the child process
    # inherits a legacy console code page. JSON.parse restores the characters.
    print(json.dumps({"fileName": output.name, "filePath": str(output)}, ensure_ascii=True))
    return 0


def find_product_sheet(
    worksheets: list[Worksheet],
) -> tuple[Worksheet, int, dict[str, int]]:
    for worksheet in worksheets:
        for row_index in range(1, min(worksheet.max_row, 50) + 1):
            columns = locate_headers(worksheet, row_index)
            if len(columns) == len(HEADER_PATTERNS):
                return worksheet, row_index, columns
    required = "발주번호, 물류센터, 입고유형, 입고예정일, SKU, 바코드, 상품명, 확정수량, 송장번호, 납품수량"
    raise ValueError(f"상품목록에서 쓰기에 필요한 헤더를 찾을 수 없습니다: {required}")


def locate_headers(worksheet: Worksheet, row_index: int) -> dict[str, int]:
    result: dict[str, int] = {}
    for cell in worksheet[row_index]:
        normalized = normalize_header(cell.value)
        if not normalized:
            continue
        for field, patterns in HEADER_PATTERNS.items():
            if field not in result and any(pattern in normalized for pattern in patterns):
                result[field] = cell.column
    return result


def write_carton_rows(
    worksheet: Worksheet,
    header_row: int,
    columns: dict[str, int],
    manifest: dict[str, Any],
) -> None:
    data_rows = [
        row_index
        for row_index in range(header_row + 1, worksheet.max_row + 1)
        if text(worksheet.cell(row_index, columns["order_no"]).value)
        and text(worksheet.cell(row_index, columns["sku_code"]).value)
    ]
    if not data_rows:
        raise ValueError("상품목록에 복제할 발주 행이 없습니다.")

    max_column = worksheet.max_column
    templates: dict[tuple[str, str], RowTemplate] = {}
    for row_index in data_rows:
        key = (
            text(worksheet.cell(row_index, columns["order_no"]).value),
            text(worksheet.cell(row_index, columns["sku_code"]).value),
        )
        templates.setdefault(key, snapshot_row(worksheet, row_index, max_column))

    order_items: dict[tuple[str, str], dict[str, Any]] = {}
    for order in manifest["orders"]:
        order_no = text(order.get("orderNo"))
        for item in order.get("items", []):
            sku_code = text(item.get("skuCode"))
            order_items[(order_no, sku_code)] = {
                "center_code": text(order.get("centerCode")),
                "transport_type": text(order.get("transportType")) or "쉽먼트",
                "expected_inbound_date": compact_date(order.get("expectedInboundDate")),
                "sku_name": text(item.get("skuName")),
                "barcode": text(item.get("barcode")),
                "confirmed_quantity": positive_int(
                    item.get("confirmedQuantity"), "confirmedQuantity"
                ),
            }
    cartons_by_batch: dict[str, list[dict[str, Any]]] = {}
    for carton in manifest["cartons"]:
        cartons_by_batch.setdefault(text(carton.get("batchId")), []).append(carton)

    output_rows: list[tuple[RowTemplate, dict[str, Any]]] = []
    seen_keys: set[tuple[str, str]] = set()
    fallback_template = snapshot_row(worksheet, data_rows[0], max_column)
    for batch in manifest["batches"]:
        batch_id = text(batch.get("id"))
        order_no = text(batch.get("orderNo"))
        sku_code = text(batch.get("skuCode"))
        key = (order_no, sku_code)
        if not batch_id or not order_no or not sku_code:
            raise ValueError("batch에는 id, orderNo, skuCode가 필요합니다.")
        if key in seen_keys:
            raise ValueError(f"중복된 PO+SKU 배치입니다: {order_no}/{sku_code}")
        seen_keys.add(key)
        order_item = order_items.get(key)
        if order_item is None:
            raise ValueError(f"orders manifest에 없는 PO+SKU입니다: {order_no}/{sku_code}")
        template = templates.get(key, fallback_template)

        carton_count = positive_int(batch.get("cartonCount"), "cartonCount")
        cartons = sorted(cartons_by_batch.get(batch_id, []), key=lambda item: int(item["cartonIndex"]))
        if len(cartons) != carton_count:
            raise ValueError(
                f"{order_no}/{sku_code}의 카톤 수가 일치하지 않습니다: "
                f"batch={carton_count}, cartons={len(cartons)}"
            )
        expected_indexes = list(range(1, carton_count + 1))
        actual_indexes = [positive_int(item.get("cartonIndex"), "cartonIndex") for item in cartons]
        if actual_indexes != expected_indexes:
            raise ValueError(f"{order_no}/{sku_code}의 cartonIndex는 1부터 연속이어야 합니다.")
        for carton in cartons:
            slip_no = text(carton.get("slipNo"))
            if not slip_no:
                raise ValueError(f"{order_no}/{sku_code} 카톤에 송장번호가 없습니다.")
            output_rows.append(
                (
                    template,
                    {
                        "order_no": order_no,
                        "center_code": order_item["center_code"],
                        "transport_type": order_item["transport_type"],
                        "expected_inbound_date": order_item["expected_inbound_date"],
                        "sku_code": sku_code,
                        "barcode": order_item["barcode"],
                        "sku_name": order_item["sku_name"],
                        "confirmed_quantity": order_item["confirmed_quantity"],
                        "slip_no": slip_no,
                        "quantity": positive_int(carton.get("quantity"), "quantity"),
                    },
                )
            )

    first_data_row = min(data_rows)
    last_data_row = max(data_rows)
    worksheet.delete_rows(first_data_row, last_data_row - first_data_row + 1)
    worksheet.insert_rows(first_data_row, len(output_rows))
    for offset, (template, values) in enumerate(output_rows):
        target_row = first_data_row + offset
        restore_row(worksheet, target_row, template)
        for field, value in values.items():
            worksheet.cell(target_row, columns[field]).value = value

    if worksheet.auto_filter.ref:
        start = worksheet.auto_filter.ref.split(":", 1)[0]
        end_column = worksheet.cell(header_row, max_column).column_letter
        worksheet.auto_filter.ref = f"{start}:{end_column}{first_data_row + len(output_rows) - 1}"


def snapshot_row(worksheet: Worksheet, row_index: int, max_column: int) -> RowTemplate:
    cells: list[dict[str, Any]] = []
    for column in range(1, max_column + 1):
        cell = worksheet.cell(row_index, column)
        cells.append(
            {
                "value": cell.value,
                "style": copy(cell._style),
                "hyperlink": copy(cell.hyperlink),
                "comment": copy(cell.comment),
            }
        )
    dimension = worksheet.row_dimensions[row_index]
    return RowTemplate(
        cells=tuple(cells),
        height=dimension.height,
        hidden=bool(dimension.hidden),
        outline_level=dimension.outlineLevel,
    )


def restore_row(worksheet: Worksheet, row_index: int, template: RowTemplate) -> None:
    for column, saved in enumerate(template.cells, start=1):
        cell: Cell = worksheet.cell(row_index, column)
        cell.value = saved["value"]
        cell._style = copy(saved["style"])
        cell.hyperlink = copy(saved["hyperlink"])
        cell.comment = copy(saved["comment"])
    dimension = worksheet.row_dimensions[row_index]
    dimension.height = template.height
    dimension.hidden = template.hidden
    dimension.outlineLevel = template.outline_level


def normalize_header(value: Any) -> str:
    return re.sub(r"[^0-9a-z가-힣]+", "", text(value).lower())


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def compact_date(value: Any) -> str:
    normalized = re.sub(r"[^0-9]", "", text(value))
    if len(normalized) != 8:
        raise ValueError(f"입고예정일은 YYYY-MM-DD 또는 YYYYMMDD여야 합니다: {value}")
    return normalized


def positive_int(value: Any, label: str) -> int:
    try:
        result = int(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"{label}은 1 이상의 정수여야 합니다.") from error
    if result <= 0:
        raise ValueError(f"{label}은 1 이상의 정수여야 합니다.")
    return result


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:  # Keep the TypeScript adapter error compact.
        print(str(error), file=sys.stderr)
        raise SystemExit(1)
