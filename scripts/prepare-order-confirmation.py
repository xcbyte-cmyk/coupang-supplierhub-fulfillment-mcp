#!/usr/bin/env python3
"""Copy a Supplier Hub confirmation template and mirror ordered quantity to confirmed quantity."""

from __future__ import annotations

import argparse
import json
import posixpath
import re
import shutil
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


HEADER_ALIASES = {
    "order_no": {"발주번호", "발주서번호", "purchaseordernumber", "poid"},
    "ordered_quantity": {"발주수량", "주문수량", "발주총수량", "orderedqty", "orderqty"},
    "confirmed_quantity": {"확정수량", "납품확정수량", "confirmedqty", "confirmqty"},
}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    source = Path(args.source).resolve()
    output = Path(args.output).resolve()
    if source == output:
        raise ValueError("다운로드한 원본 양식을 덮어쓸 수 없습니다.")
    if output.exists():
        raise FileExistsError(f"출력 파일이 이미 존재합니다: {output}")

    payload = json.load(sys.stdin)
    order_nos = {text(value) for value in payload.get("orderNos", []) if text(value)}
    if not order_nos:
        raise ValueError("확정할 발주번호가 없습니다.")

    workbook = load_workbook(source, read_only=False, data_only=False)
    worksheet, header_row, columns = find_sheet(workbook.worksheets)
    changed_rows, cell_edits = plan_quantity_edits(
        worksheet, header_row, columns, order_nos
    )
    worksheet_path = worksheet_archive_path(source, worksheet.title)
    workbook.close()
    output.parent.mkdir(parents=True, exist_ok=True)
    if cell_edits:
        write_patched_workbook(source, output, worksheet_path, cell_edits)
    else:
        shutil.copy2(source, output)
    print(
        json.dumps(
            {
                "fileName": output.name,
                "filePath": str(output),
                "orderNos": sorted(order_nos),
                "changedRows": changed_rows,
            },
            ensure_ascii=True,
        )
    )
    return 0


def find_sheet(
    worksheets: list[Worksheet],
) -> tuple[Worksheet, int, dict[str, int]]:
    for worksheet in worksheets:
        for row_index in range(1, min(worksheet.max_row, 50) + 1):
            columns = locate_headers(worksheet, row_index)
            if len(columns) == len(HEADER_ALIASES):
                return worksheet, row_index, columns
    raise ValueError("확정 엑셀에 필요한 발주번호, 발주수량, 확정수량 헤더를 찾을 수 없습니다.")


def locate_headers(worksheet: Worksheet, row_index: int) -> dict[str, int]:
    result: dict[str, int] = {}
    for cell in worksheet[row_index]:
        normalized = normalize_header(cell.value)
        for field, aliases in HEADER_ALIASES.items():
            if field not in result and normalized in aliases:
                result[field] = cell.column
    return result


def plan_quantity_edits(
    worksheet: Worksheet,
    header_row: int,
    columns: dict[str, int],
    order_nos: set[str],
) -> tuple[int, list[tuple[str, str]]]:
    found: set[str] = set()
    changed_rows = 0
    cell_edits: list[tuple[str, str]] = []
    for row_index in range(header_row + 1, worksheet.max_row + 1):
        order_no = text(worksheet.cell(row_index, columns["order_no"]).value)
        if order_no not in order_nos:
            continue
        ordered = worksheet.cell(row_index, columns["ordered_quantity"]).value
        if ordered is None or text(ordered) == "":
            raise ValueError(f"발주 {order_no}의 발주수량이 비어 있습니다.")
        ordered_cell = worksheet.cell(row_index, columns["ordered_quantity"])
        confirmed_cell = worksheet.cell(row_index, columns["confirmed_quantity"])
        if confirmed_cell.value != ordered_cell.value:
            cell_edits.append((ordered_cell.coordinate, confirmed_cell.coordinate))
        found.add(order_no)
        changed_rows += 1
    missing = sorted(order_nos - found)
    if missing:
        raise ValueError(f"확정 양식에서 선택한 발주를 찾을 수 없습니다: {', '.join(missing)}")
    return changed_rows, cell_edits


def write_patched_workbook(
    source: Path,
    output: Path,
    worksheet_path: str,
    cell_edits: list[tuple[str, str]],
) -> None:
    """Patch only target cell XML so unrelated workbook cells and styles stay byte-stable."""
    with ZipFile(source, "r") as input_zip, ZipFile(output, "w") as output_zip:
        patched = False
        for info in input_zip.infolist():
            data = input_zip.read(info.filename)
            if info.filename == worksheet_path:
                patched = True
                xml = data.decode("utf-8")
                for source_coordinate, target_coordinate in cell_edits:
                    xml = copy_cell_xml_value(xml, source_coordinate, target_coordinate)
                data = xml.encode("utf-8")
            output_zip.writestr(info, data)
    if not patched:
        output.unlink(missing_ok=True)
        raise ValueError(f"확정 양식의 워크시트 파일을 찾을 수 없습니다: {worksheet_path}")


def worksheet_archive_path(source: Path, worksheet_title: str) -> str:
    spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    office_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    with ZipFile(source, "r") as workbook_zip:
        workbook_root = ET.fromstring(workbook_zip.read("xl/workbook.xml"))
        relationship_id = None
        for sheet in workbook_root.findall(f".//{{{spreadsheet_ns}}}sheet"):
            if sheet.get("name") == worksheet_title:
                relationship_id = sheet.get(f"{{{office_rel_ns}}}id")
                break
        if not relationship_id:
            raise ValueError(f"확정 양식에서 워크시트를 찾을 수 없습니다: {worksheet_title}")
        rels_root = ET.fromstring(workbook_zip.read("xl/_rels/workbook.xml.rels"))
        target = None
        for relationship in rels_root.findall(f"{{{package_rel_ns}}}Relationship"):
            if relationship.get("Id") == relationship_id:
                target = relationship.get("Target")
                break
        if not target:
            raise ValueError(f"확정 양식의 워크시트 연결을 찾을 수 없습니다: {worksheet_title}")
        if target.startswith("/"):
            return target.lstrip("/")
        return posixpath.normpath(posixpath.join("xl", target))


def copy_cell_xml_value(xml: str, source_coordinate: str, target_coordinate: str) -> str:
    source_match = find_cell_xml(xml, source_coordinate)
    target_match = find_cell_xml(xml, target_coordinate)
    source_open, source_body = source_match.group(1), source_match.group(2)
    target_open = target_match.group(1)

    replacement_open = re.sub(
        rf'\br="{re.escape(source_coordinate)}"',
        f'r="{target_coordinate}"',
        source_open,
        count=1,
    )
    target_style = re.search(r'\bs="([^"]*)"', target_open)
    if target_style:
        if re.search(r'\bs="[^"]*"', replacement_open):
            replacement_open = re.sub(
                r'\bs="[^"]*"', f's="{target_style.group(1)}"', replacement_open, count=1
            )
        else:
            replacement_open = replacement_open[:-1] + f' s="{target_style.group(1)}">'
    else:
        replacement_open = re.sub(r'\s+s="[^"]*"', "", replacement_open, count=1)

    replacement = f"{replacement_open}{source_body}</c>"
    return xml[: target_match.start()] + replacement + xml[target_match.end() :]


def find_cell_xml(xml: str, coordinate: str) -> re.Match[str]:
    pattern = re.compile(
        rf'(<c\b[^>]*\br="{re.escape(coordinate)}"[^>]*>)(.*?)(</c>)',
        re.DOTALL,
    )
    matches = list(pattern.finditer(xml))
    if len(matches) != 1:
        raise ValueError(
            f"확정 양식의 셀 {coordinate}을(를) 하나로 식별할 수 없습니다."
        )
    return matches[0]


def normalize_header(value: Any) -> str:
    return re.sub(r"[^0-9a-z가-힣]+", "", text(value).lower())


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(str(error), file=sys.stderr)
        raise SystemExit(1)
